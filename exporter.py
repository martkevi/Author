"""Manuscript export: PDF (6x9 in, KDP trim), DOCX, and EPUB."""
import re
import uuid
import zipfile
from datetime import datetime
from xml.sax.saxutils import escape

import config
import storage


def _slug(title: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", title).strip("_")[:50] or "novel"


def _paras(text: str) -> list:
    out = []
    for p in re.split(r"\n\s*\n", text.strip()):
        p = " ".join(p.split())
        if p:
            out.append(p)
    return out


def _is_break(p: str) -> bool:
    return len(p.replace(" ", "")) >= 1 and re.fullmatch(r"[\*\-#•~\s]+", p) is not None


def _italics(p: str) -> list:
    """Split a paragraph into (text, is_italic) parts using *asterisks*."""
    parts = []
    for seg in re.split(r"(\*[^*\n]+\*)", p):
        if len(seg) > 2 and seg.startswith("*") and seg.endswith("*"):
            parts.append((seg[1:-1], True))
        elif seg:
            parts.append((seg.replace("*", ""), False))
    return parts


def _load(book_id: int):
    book = storage.get_book(book_id)
    chapters = storage.get_chapters(book_id)
    if not chapters:
        raise ValueError("No chapters written yet")
    return book, chapters


def _byline() -> str:
    return "A Novel"


# ---- PDF ------------------------------------------------------------------------
def export_pdf(book_id: int) -> str:
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.pagesizes import inch
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    book, chapters = _load(book_id)
    path = config.EXPORT_DIR / f"{_slug(book['title'])}.pdf"

    def safe(s: str) -> str:  # standard PDF fonts use cp1252
        s = s.replace("\u2011", "-").replace("\u2060", "")
        return s.encode("cp1252", "replace").decode("cp1252")

    def markup(p: str) -> str:
        return "".join((f"<i>{escape(safe(t))}</i>" if it else escape(safe(t))) for t, it in _italics(p))

    body = ParagraphStyle("body", fontName="Times-Roman", fontSize=11, leading=15.5,
                          alignment=TA_JUSTIFY, firstLineIndent=16, spaceAfter=0)
    first = ParagraphStyle("first", parent=body, firstLineIndent=0)
    brk = ParagraphStyle("brk", parent=body, alignment=TA_CENTER, firstLineIndent=0, spaceBefore=10, spaceAfter=10)
    title_s = ParagraphStyle("t", fontName="Times-Bold", fontSize=30, leading=36, alignment=TA_CENTER)
    sub_s = ParagraphStyle("s", fontName="Times-Italic", fontSize=14, leading=18, alignment=TA_CENTER)
    chnum = ParagraphStyle("cn", fontName="Times-Roman", fontSize=11, leading=14, alignment=TA_CENTER, textColor="#555555")
    chtitle = ParagraphStyle("ct", fontName="Times-Bold", fontSize=20, leading=24, alignment=TA_CENTER, spaceAfter=22)
    small = ParagraphStyle("sm", parent=body, fontSize=9, leading=12, alignment=TA_CENTER, firstLineIndent=0)

    def footer(canvas, doc):
        if doc.page > 2:
            canvas.saveState()
            canvas.setFont("Times-Roman", 9)
            canvas.drawCentredString(3 * inch, 0.45 * inch, str(doc.page - 2))
            canvas.restoreState()

    story = [Spacer(1, 2.2 * inch), Paragraph(escape(safe(book["title"])), title_s)]
    sub = book["data"].get("bible", {}).get("subtitle")
    if sub:
        story += [Spacer(1, 14), Paragraph(escape(safe(str(sub))), sub_s)]
    story += [Spacer(1, 30), Paragraph(_byline(), sub_s), PageBreak(),
              Spacer(1, 4 * inch),
              Paragraph(f"Copyright \u00a9 {datetime.now().year}. All rights reserved.", small), Spacer(1, 8),
              Paragraph("This is a work of fiction. Names, characters, places, and events are products of the "
                        "author's imagination or are used fictitiously. Any resemblance to actual persons, "
                        "living or dead, is coincidental.", small)]
    for c in chapters:
        story += [PageBreak(), Spacer(1, 0.9 * inch), Paragraph(f"CHAPTER {c['n'] + 1}", chnum),
                  Paragraph(escape(safe(c["title"])), chtitle)]
        use_first = True
        for p in _paras(c["text"]):
            if _is_break(p):
                story.append(Paragraph("*&nbsp;&nbsp;&nbsp;*&nbsp;&nbsp;&nbsp;*", brk))
                use_first = True
            else:
                story.append(Paragraph(markup(p), first if use_first else body))
                use_first = False

    doc = SimpleDocTemplate(str(path), pagesize=(6 * inch, 9 * inch), leftMargin=0.8 * inch,
                            rightMargin=0.7 * inch, topMargin=0.8 * inch, bottomMargin=0.85 * inch,
                            title=book["title"], author="")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return str(path)


# ---- DOCX -----------------------------------------------------------------------
def export_docx(book_id: int) -> str:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    book, chapters = _load(book_id)
    path = config.EXPORT_DIR / f"{_slug(book['title'])}.docx"
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(6), Inches(9)
    sec.left_margin, sec.right_margin = Inches(0.85), Inches(0.75)
    sec.top_margin, sec.bottom_margin = Inches(0.8), Inches(0.85)

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(11.5)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(0)
    h1 = doc.styles["Heading 1"]
    h1.font.name, h1.font.size, h1.font.bold = "Times New Roman", Pt(20), True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before, h1.paragraph_format.space_after = Pt(60), Pt(24)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(150)
    r = t.add_run(book["title"])
    r.bold, r.font.size = True, Pt(30)
    sub = book["data"].get("bible", {}).get("subtitle")
    if sub:
        s = doc.add_paragraph()
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        s.add_run(str(sub)).italic = True
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.page_break_before = True
    c.paragraph_format.space_before = Pt(240)
    c.add_run(f"Copyright \u00a9 {datetime.now().year}. All rights reserved.\n\nThis is a work of fiction. "
              "Names, characters, places, and events are products of the author's imagination or are used "
              "fictitiously. Any resemblance to actual persons, living or dead, is coincidental.").font.size = Pt(9)

    for ch in chapters:
        h = doc.add_heading(f"Chapter {ch['n'] + 1}: {ch['title']}", level=1)
        h.paragraph_format.page_break_before = True
        first = True
        for p in _paras(ch["text"]):
            para = doc.add_paragraph()
            if _is_break(p):
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.add_run("*   *   *")
                first = True
                continue
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if not first:
                para.paragraph_format.first_line_indent = Inches(0.3)
            first = False
            for txt, it in _italics(p):
                para.add_run(txt).italic = it
    doc.core_properties.title = book["title"]
    doc.save(str(path))
    return str(path)


# ---- EPUB -----------------------------------------------------------------------
_CSS = """body{font-family:serif;line-height:1.5;margin:5%}
h1{text-align:center;margin:3em 0 1.2em;font-size:1.6em}
p{margin:0;text-indent:1.3em;text-align:justify}
p.first,p.brk{text-indent:0}
p.brk{text-align:center;margin:1.2em 0}
.title{text-align:center;margin-top:30%}
.title h1{font-size:2.2em}
.small{font-size:.8em;text-indent:0;text-align:center;margin-top:3em}"""


def export_epub(book_id: int) -> str:
    book, chapters = _load(book_id)
    path = config.EXPORT_DIR / f"{_slug(book['title'])}.epub"
    uid = f"urn:uuid:{uuid.uuid4()}"
    title = escape(book["title"])
    xhtml = ('<?xml version="1.0" encoding="utf-8"?><!DOCTYPE html><html xmlns="http://www.w3.org/1999/xhtml" '
             'xmlns:epub="http://www.idpf.org/2007/ops" lang="en"><head><title>{t}</title>'
             '<link rel="stylesheet" type="text/css" href="style.css"/></head><body>{b}</body></html>')

    def md(p: str) -> str:
        return "".join((f"<em>{escape(t)}</em>" if it else escape(t)) for t, it in _italics(p))

    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml",
                   '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                   '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
                   '</rootfiles></container>', compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("OEBPS/style.css", _CSS, compress_type=zipfile.ZIP_DEFLATED)
        sub = book["data"].get("bible", {}).get("subtitle")
        title_body = (f'<div class="title"><h1>{title}</h1>' + (f"<p class='brk'><em>{escape(str(sub))}</em></p>" if sub else "")
                      + f'<p class="small">Copyright &#169; {datetime.now().year}. All rights reserved. This is a work of fiction. '
                        'Names, characters, places, and events are products of the author\'s imagination or are used fictitiously.</p></div>')
        z.writestr("OEBPS/title.xhtml", xhtml.format(t=title, b=title_body), compress_type=zipfile.ZIP_DEFLATED)

        items, spine, nav_li = ['<item id="title" href="title.xhtml" media-type="application/xhtml+xml"/>'], ['<itemref idref="title"/>'], []
        for c in chapters:
            fn, cid = f"ch{c['n'] + 1}.xhtml", f"ch{c['n'] + 1}"
            heading = f"Chapter {c['n'] + 1}: {escape(c['title'])}"
            body, first = [f"<h1>{heading}</h1>"], True
            for p in _paras(c["text"]):
                if _is_break(p):
                    body.append('<p class="brk">*&#160;&#160;&#160;*&#160;&#160;&#160;*</p>')
                    first = True
                else:
                    body.append(f'<p class="first">{md(p)}</p>' if first else f"<p>{md(p)}</p>")
                    first = False
            z.writestr(f"OEBPS/{fn}", xhtml.format(t=heading, b="".join(body)), compress_type=zipfile.ZIP_DEFLATED)
            items.append(f'<item id="{cid}" href="{fn}" media-type="application/xhtml+xml"/>')
            spine.append(f'<itemref idref="{cid}"/>')
            nav_li.append(f'<li><a href="{fn}">{heading}</a></li>')

        nav = xhtml.format(t="Contents", b='<nav epub:type="toc"><h1>Contents</h1><ol>' + "".join(nav_li) + "</ol></nav>")
        z.writestr("OEBPS/nav.xhtml", nav, compress_type=zipfile.ZIP_DEFLATED)
        opf = ('<?xml version="1.0" encoding="utf-8"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" '
               'unique-identifier="bid"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
               f'<dc:identifier id="bid">{uid}</dc:identifier><dc:title>{title}</dc:title><dc:language>en</dc:language>'
               f'<meta property="dcterms:modified">{datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ")}</meta></metadata>'
               '<manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>'
               '<item id="css" href="style.css" media-type="text/css"/>' + "".join(items) +
               '</manifest><spine>' + "".join(spine) + '</spine></package>')
        z.writestr("OEBPS/content.opf", opf, compress_type=zipfile.ZIP_DEFLATED)
    return str(path)


EXPORTERS = {"pdf": export_pdf, "docx": export_docx, "epub": export_epub}
